from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = ROOT / "doc"
SHOT_DIR = DOC_DIR / "assets" / "操作手册V2截图"
OUTPUT = DOC_DIR / "高校实践管理系统操作手册-V2.0.docx"

BLUE = "2563EB"
BLUE_DARK = "1E3A8A"
BLUE_LIGHT = "EFF6FF"
GRAY = "64748B"
GRAY_LIGHT = "F8FAFC"
GREEN_LIGHT = "ECFDF5"
AMBER_LIGHT = "FFFBEB"
RED_LIGHT = "FEF2F2"
FONT = "Noto Sans SC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_run_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in (
        ("Title", 26, BLUE_DARK),
        ("Heading 1", 17, BLUE_DARK),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11, "334155"),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 16)
        style.paragraph_format.space_after = Pt(6)

    caption = styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(GRAY)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("成都锦城学院  高校实践管理系统操作手册 V2.0")
    set_run_font(run, size=8.5, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_run_font(run, size=8.5, color=GRAY)
    add_field(footer, "PAGE")
    run = footer.add_run(" 页")
    set_run_font(run, size=8.5, color=GRAY)

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(document):
    for _ in range(4):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("成都锦城学院")
    set_run_font(run, size=18, bold=True, color=BLUE_DARK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("高校实践管理系统")
    set_run_font(run, size=30, bold=True, color=BLUE_DARK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run("交付操作手册")
    set_run_font(run, size=24, bold=True, color=BLUE)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    run = p.add_run("适用终端：PC 管理端 / H5 移动端")
    set_run_font(run, size=11, color=GRAY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("文档版本：V2.0（完整交付版）")
    set_run_font(run, size=11, color=GRAY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("编制日期：2026 年 8 月 11 日")
    set_run_font(run, size=11, color=GRAY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(50)
    run = p.add_run("本手册基于当前本地交付版本和实际页面编制")
    set_run_font(run, size=9.5, color=GRAY)
    document.add_page_break()


def add_heading(document, text, level=1):
    return document.add_heading(text, level=level)


def add_body(document, text, bold_prefix=None):
    p = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        run = p.add_run(item)
        set_run_font(run)


def add_steps(document, items):
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        run = p.add_run(item)
        set_run_font(run)


def add_note(document, text, kind="note"):
    fills = {"note": BLUE_LIGHT, "warning": AMBER_LIGHT, "danger": RED_LIGHT, "success": GREEN_LIGHT}
    labels = {"note": "说明", "warning": "注意", "danger": "重要", "success": "建议"}
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fills[kind])
    set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
    p = cell.paragraphs[0]
    run = p.add_run(f"{labels[kind]}：")
    set_run_font(run, size=9.5, bold=True, color=BLUE_DARK if kind == "note" else "92400E")
    run = p.add_run(text)
    set_run_font(run, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(value))
        set_run_font(run, size=9, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, GRAY_LIGHT)
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            set_run_font(run, size=8.8)
            if widths:
                cell.width = Cm(widths[col_index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(document, filename, caption, width=16.3):
    path = SHOT_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    cp = document.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = False
    cp.add_run(caption)


def add_mobile_pair(document, left, left_caption, right, right_caption):
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, filename in enumerate((left, right)):
        path = SHOT_DIR / filename
        if not path.exists():
            raise FileNotFoundError(path)
        cell = table.cell(0, index)
        set_cell_margins(cell, top=40, start=80, bottom=40, end=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(6.7))
    for index, caption in enumerate((left_caption, right_caption)):
        cell = table.cell(1, index)
        set_cell_margins(cell, top=20, start=80, bottom=100, end=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(caption)
        set_run_font(run, size=8.8, color=GRAY)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_toc(document):
    document.add_page_break()
    add_heading(document, "目录", 1)
    add_table(
        document,
        ["章节", "内容"],
        [
            ["1", "文档说明"],
            ["2", "系统访问与登录"],
            ["3", "PC 工作台"],
            ["4", "用户、基础档案与学校配置"],
            ["5", "实习管理"],
            ["6", "实验实训管理"],
            ["7", "社会实践管理"],
            ["8", "消息中心与待办模板"],
            ["9", "文件、文档、模板、导出与收藏"],
            ["10", "统计报表与日志审计"],
            ["11", "H5 移动端通用操作"],
            ["12", "常见问题"],
            ["13", "日常维护顺序"],
            ["14", "附录"],
        ],
        [2.5, 14.0],
    )
    document.add_page_break()


def build_manual():
    document = Document()
    configure_document(document)
    add_cover(document)

    add_heading(document, "文档控制", 1)
    add_table(
        document,
        ["项目", "内容"],
        [
            ["文档名称", "高校实践管理系统交付操作手册"],
            ["当前版本", "V2.0"],
            ["适用范围", "成都锦城学院 PC 管理端与 H5 移动端"],
            ["覆盖角色", "超级管理员、学校管理员、学院管理员、专业管理员、教师、学生"],
            ["编制依据", "当前本地交付版本、实际菜单权限和页面操作"],
        ],
        [3.5, 13.0],
    )
    add_heading(document, "修订记录", 2)
    add_table(
        document,
        ["版本", "日期", "修订内容"],
        [
            ["V1.0", "2026-07-14", "首版交付手册"],
            ["V2.0", "2026-08-11", "按最新桌面、实习计划任务、实习实施、实验实训一体化、社会实践、消息模板和 H5 角色界面重写"],
        ],
        [2.0, 3.0, 11.5],
    )
    add_note(document, "截图来自当前本地环境，页面中的人员、时间、任务和业务记录为演示数据；正式使用时以学校实际数据为准。")
    add_toc(document)

    add_heading(document, "1 文档说明", 1)
    add_heading(document, "1.1 使用对象", 2)
    add_body(document, "本手册用于系统交付、业务培训和日常操作查询。不同角色仅能看到本人菜单权限、按钮权限和数据范围内的内容。")
    add_table(
        document,
        ["角色", "主要职责", "默认数据范围"],
        [
            ["超级管理员", "全局配置、学校配置、权限、消息模板、数据维护", "全校"],
            ["学校管理员", "学校业务配置、计划审核、数据查看、用户与基础档案维护", "全校"],
            ["学院管理员", "学院计划、任务、基地、审核与统计", "授权学院"],
            ["专业管理员", "专业计划、任务、学生范围和审核", "授权专业"],
            ["教师", "执行本人任务、审核学生材料、评分和提交教师材料", "本人负责的任务与学生"],
            ["学生", "查看本人任务、签到、提交材料、查看审核记录和成绩", "本人数据"],
        ],
        [2.5, 8.0, 6.0],
    )

    add_heading(document, "1.2 业务维度", 2)
    add_bullets(document, [
        "系统按学校独立业务库运行，业务页面只显示学校数据和设置。",
        "毕业实习按毕业届次管理；其他实习、实验、实训和社会实践按年级管理。",
        "基础档案顺序为年级或毕业届次、学院、专业、班级、用户。",
        "筛选条件默认应用当前年级或届次，并根据当前账号自动限定学院、专业或本人任务范围。",
        "学生账号完成学籍或任务绑定后才能进入对应业务。",
    ])

    add_heading(document, "1.3 状态与操作规则", 2)
    add_table(
        document,
        ["状态", "含义", "可执行操作"],
        [
            ["草稿", "内容已保存但未提交", "继续编辑、提交审核"],
            ["待审核", "已提交给当前审核节点", "审核人通过或退回"],
            ["需修改", "审核后要求提交人修改", "提交人修改并重新提交"],
            ["已通过", "当前流程审核完成", "查看、记录；有权限者可发起通过后修改"],
            ["启用/停用", "基础档案或配置是否可用", "通过开关切换"],
        ],
        [2.4, 7.2, 6.9],
    )
    add_bullets(document, [
        "提交人使用“保存草稿”和“提交审核”控制状态，不直接选择业务状态。",
        "审核人可先保存评阅草稿；提交审核后才发送待办或处理结果消息。",
        "已通过记录不能再次普通通过或退回，只能按权限执行“通过后修改”。",
        "每次提交形成 recording 记录，每次审核形成 review_opinion 记录，流程记录按提交主线和审核支线展示。",
        "提交期间按钮自动禁用，服务端同时使用事务、状态校验和 Redis 锁避免重复或并发处理。",
    ])

    add_heading(document, "2 系统访问与登录", 1)
    add_heading(document, "2.1 访问入口", 2)
    add_bullets(document, [
        "统一入口：访问学校部署域名根路径，系统根据浏览器头和屏幕尺寸进入 PC 或 H5，默认进入 PC。",
        "PC 直接入口：/pc/index.html。",
        "H5 直接入口：/h5/index.html。",
        "本地访问地址：http://127.0.0.1:8787/。",
    ])
    add_heading(document, "2.2 账号密码登录", 2)
    add_steps(document, [
        "输入管理员分配的登录账号。",
        "输入密码并点击“登录”。",
        "系统加载当前身份、角色权限、菜单权限和数据范围。",
        "首次使用后按学校要求更新个人资料和密码。",
    ])
    add_figure(document, "01-PC登录页.png", "图 1  PC 端登录页")
    add_note(document, "登录页不显示默认账号和密码；测试账号不得写入正式交付材料。", "warning")

    add_heading(document, "2.3 多账号切换与一键登录", 2)
    add_bullets(document, [
        "同一自然人绑定多个账号时，个人菜单才显示“切换身份”；没有可切换账号时不显示入口。",
        "切换身份后系统重新加载角色、菜单和数据范围，不需要重复输入密码。",
        "用户管理中的“一键登录”按管理员数据权限开放。链接有效期 600 秒、一次有效，未过期前重复点击复用同一链接。",
        "超级管理员可生成全校账号链接；学校管理员可操作全校；学院和专业管理员只能操作授权范围内的教师和学生。",
    ])

    add_heading(document, "3 PC 工作台", 1)
    add_heading(document, "3.1 桌面、窗口与任务栏", 2)
    add_bullets(document, [
        "桌面显示固定模块和用户添加的快捷方式；默认固定模块不能移除。",
        "业务窗口支持拖动、调整大小、最小化、最大化和关闭，也可同时打开多个模块。",
        "任务栏集中显示启动台、全局搜索、运行窗口、消息入口、在线状态和时间。",
        "点击已打开窗口的任务栏图标可聚焦或最小化，不会重复生成同一模块窗口。",
        "浏览器非全屏时，桌面和任务栏会按可用宽度自适应。",
    ])
    add_figure(document, "02-PC桌面与用户管理.png", "图 2  PC 工作台、窗口和固定操作列")

    add_heading(document, "3.2 启动台与快捷方式", 2)
    add_steps(document, [
        "点击任务栏的启动台图标。",
        "在顶部搜索框输入模块名称，或按教学业务、数据与配置、个人工具分类查找。",
        "点击模块图标打开模块。",
        "点击模块右上角加号加入桌面；已添加后显示勾选状态。",
        "点击空白区域或关闭按钮退出启动台。",
    ])
    add_figure(document, "03-PC启动台.png", "图 3  PC 启动台与模块快捷方式")

    add_heading(document, "3.3 全局搜索与收藏夹", 2)
    add_bullets(document, [
        "任务栏搜索支持按模块名称和关键词定位可见模块，结果受菜单权限控制。",
        "收藏夹用于维护常用网址、标题和图标；打开收藏时在新页面访问外部网址。",
        "收藏较多时可按关键词搜索，也可将收藏项添加到桌面。",
        "仅 PC 端提供收藏夹和桌面快捷方式。",
    ])

    add_heading(document, "4 用户、基础档案与学校配置", 1)
    add_heading(document, "4.1 用户管理", 2)
    add_steps(document, [
        "从桌面或启动台打开“用户管理”。",
        "使用姓名、账号、手机号、角色或状态查询用户。",
        "点击“新增用户”，录入基本资料、登录账号、初始密码和角色；新增教师、学生或管理员即完成账号注册。",
        "按需执行编辑、启用或停用、重置密码、一键登录、查看日志和查看绑定账号。",
        "用户日志采用分页展示，操作列固定在列表右侧。",
    ])
    add_figure(document, "02-PC桌面与用户管理.png", "图 4  用户管理列表")
    add_note(document, "新增学生时应同时维护年级或毕业届次、学院、专业、班级；教师和管理角色应配置组织范围。", "warning")

    add_heading(document, "4.2 年级、毕业届次、学院、专业和班级", 2)
    add_table(
        document,
        ["档案", "用途", "维护要点"],
        [
            ["年级", "普通实习、实验、实训、社会实践", "可设置当前年级，作为默认筛选"],
            ["毕业届次", "毕业实习", "独立维护当前毕业届次"],
            ["学院", "组织与数据权限", "维护全称、简称和状态"],
            ["专业", "计划与任务范围", "关联学院，维护简称和状态"],
            ["班级", "任务绑定和学生范围", "关联年级、学院、专业，维护简称"],
        ],
        [2.5, 6.0, 8.0],
    )
    add_bullets(document, [
        "学院、专业、班级下拉项按上级选择逐级加载。",
        "专业和班级支持 Excel 导入；表格填写年级、学院、专业、班级，已存在记录不会重复创建。",
        "启用状态使用开关维护，关闭后新业务不再选择该档案。",
    ])

    add_heading(document, "4.3 基地管理", 2)
    add_bullets(document, [
        "基地管理维护基地名称、信用代码、联系人、联系电话、地址和启用状态。",
        "学院、专业筛选使用下拉选项，列表支持关键词和分页。",
        "基地支持 Excel 导入；导入前必须能匹配学院和专业，不能匹配的数据不写入。",
        "长期基地维护协议、场地、服务专业、服务课程、负责人、校内外指导人员等完整资料。",
        "临时基地只维护所在位置、服务专业、服务课程和基地负责人。",
        "基地建设、基地申报和基地使用均基于同一基地档案关联，避免重复建档。",
    ])
    add_figure(document, "10-PC基地管理.png", "图 5  基地管理列表")

    add_heading(document, "4.4 菜单管理", 2)
    add_bullets(document, [
        "菜单树层级为主菜单、菜单、列表、按钮。",
        "按钮必须挂在对应列表节点下，例如：实习管理 → 计划表 → 列表 → 新增计划。",
        "菜单可维护名称、权限码、平台、路由、排序、显示状态、是否模块和模块图标。",
        "设置为模块后可在启动台显示，并允许用户加入桌面。",
        "支持关键词搜索、一键展开和一键收缩；菜单名称保存后桌面与启动台同步更新。",
    ])
    add_figure(document, "15-PC菜单管理.png", "图 6  菜单树和按钮权限层级")

    add_heading(document, "4.5 角色权限与组织范围", 2)
    add_steps(document, [
        "进入“系统配置 → 角色权限”。",
        "选择需要维护的角色。",
        "按菜单树勾选主菜单、菜单、列表和按钮权限。",
        "保存后重新登录或切换身份，使权限重新加载。",
        "进入“组织范围”，为学院管理员和专业管理员配置可管理学院、专业。",
    ])
    add_figure(document, "16-PC角色权限.png", "图 7  角色菜单和按钮权限")
    add_note(document, "前端隐藏不可用入口，后端仍会独立校验权限码和组织数据范围。", "danger")

    add_heading(document, "4.6 操作说明、课节与企业微信", 2)
    add_bullets(document, [
        "每个模块右上角提供“操作说明”，内容由系统配置中的操作说明列表维护。",
        "操作说明支持富文本，可维护操作流程、常见问题和注意事项。",
        "课节配置预置第 1 至第 12 节，管理员可维护每节课开始和结束时间。",
        "企业微信应用配置包括企业和应用参数、代理地址及应用菜单，应用菜单支持可视化编辑。",
    ])

    add_heading(document, "4.7 个人设置与学校登录背景", 2)
    add_bullets(document, [
        "点击头像区域上传头像，不填写图片链接。",
        "选择预置桌面壁纸或上传自定义壁纸，保存后立即应用。",
        "超级管理员和学校管理员可上传全校 PC 登录背景。",
        "可分别开启或关闭站内消息、企业微信和邮件接收。",
        "头像和壁纸由文件服务管理，浏览器缓存壁纸信息以减少重复加载。",
    ])
    add_figure(document, "24-PC个人设置.png", "图 8  个人资料、壁纸和消息接收设置")

    add_heading(document, "4.8 数据管理", 2)
    add_bullets(document, [
        "一键清除测试数据仅在环境变量 APP_MODE=test 时允许执行。",
        "正式环境禁止执行清理操作，即使前端按钮被误触也会由后端拒绝。",
        "清理范围不包含超级管理员、学校流程、权限、消息模板和必要配置。",
        "执行前必须阅读影响范围并在二次确认框中确认。",
    ])

    add_heading(document, "5 实习管理", 1)
    add_heading(document, "5.1 完整业务流程", 2)
    add_steps(document, [
        "管理员从教务系统主动拉取或通过 Excel 导入教学计划，也可手工新增。",
        "确认教学计划后生成实习计划；计划按课程和专业维护，不直接填写学生人数。",
        "学院管理员或授权管理员在计划详情中拆分实习任务。",
        "每个任务唯一对应一名负责教师，可绑定一个或多个班级和学生。",
        "任务提交审核，通过后由任课教师进入实习实施。",
        "学生按阶段完成特殊申请、签到、日志、报告、延期申请和归档材料。",
        "教师审核本人任务学生材料并评分，管理员处理授权节点和巡查。",
        "任务信息需要调整时发起任务变更，审核通过后形成新任务记录。",
        "课程成绩按计划配置的成绩规则汇总多个任务成绩，完成归档。",
    ])
    add_figure(document, "04-PC实习管理总览.png", "图 9  实习管理总览")

    add_heading(document, "5.2 教务计划同步与计划表", 2)
    add_bullets(document, [
        "教学计划同步当前采用主动拉取方式；接口就绪后由管理员手动执行同步。",
        "同步结果先进入待确认列表，不直接生成实习计划。",
        "确认生成前需二次确认；同一来源计划按来源编号和课程信息防重复。",
        "计划保留来源课程、学分、课程性质、原始教师、时间、地点和备注。",
        "计划名称可由用户填写；未填写时按课程和专业自动生成。",
        "Excel 导入入口提供模板下载，模板按当前年级或毕业届次口径填写。",
        "计划使用“保存草稿”和“提交审核”；草稿不会发送待办，审核人不可处理。",
    ])
    add_figure(document, "05-PC计划表列表.png", "图 10  计划表列表、筛选和审核状态")

    add_heading(document, "5.3 任务拆分、审核与学生绑定", 2)
    add_steps(document, [
        "在计划表列表点击“查看”。",
        "在“任务拆分”区域点击“新增任务”。",
        "填写任务编号、任务名称、负责教师、班级、实施时间、地点和说明。",
        "系统按班级和学生绑定统计任务人数；计划本身不维护固定人数。",
        "保存草稿后可继续编辑；提交审核后进入管理员审核。",
        "审核通过后教师可执行，审核人和审核意见始终保留。",
    ])
    add_figure(document, "06-PC计划任务详情.png", "图 11  计划详情和任务拆分")
    add_note(document, "同一教师可以承担多个任务，但不同任务应明确班级、学生和时间范围，避免混淆。", "warning")

    add_heading(document, "5.4 实习实施与任务变更", 2)
    add_bullets(document, [
        "实习实施列表按年级或届次、学院、专业、计划、类型、组织方式、状态和关键词筛选。",
        "详情使用“任务概览、学生绑定、任务变更、实施表”标签切换。",
        "实施表自动带出课程、任务、教师、班级、学生、时间和地点，主要补充经费项目、金额、来源和说明。",
        "实施表可按学校样表导出 PDF；导出任务完成后在消息中心和导出中心通知。",
        "任务变更填写新教师、班级、时间、地点等调整内容及变更原因，提交审核后生效。",
        "原任务、变更申请、生效任务和审核记录均保留。",
    ])
    add_figure(document, "09-PC实习实施列表.png", "图 12  实习实施列表")
    add_figure(document, "07-PC实习实施详情.png", "图 13  实习实施任务概览")
    add_figure(document, "08-PC任务学生绑定.png", "图 14  实习实施学生绑定")

    add_heading(document, "5.5 基地建设", 2)
    add_bullets(document, [
        "超级管理员、学校管理员、学院管理员和专业管理员可按权限录入基地建设资料。",
        "长期基地录入基地基本情况、场地、服务专业和课程、基地负责人、校内外指导人员等。",
        "临时基地只录入所在地、服务专业、服务课程和负责人。",
        "学院意见在系统录入阶段留空，导出 Word 后线下签署和盖章。",
        "基地建设资料关联基地档案，避免相同基地重复录入。",
    ])

    add_heading(document, "5.6 实习大纲与实习指导书", 2)
    add_bullets(document, [
        "实习大纲和实习指导书是两种独立材料，分别维护和归档。",
        "新增时关联实习计划或任务，填写标题、版本、内容和附件。",
        "正文较长时通过“查看”进入详情，不在列表内截断后直接判断。",
        "提交审核、退回、通过后修改和流程记录均使用统一审核机制。",
    ])

    add_heading(document, "5.7 学生过程材料", 2)
    add_table(
        document,
        ["材料", "提交人", "主要内容", "处理人"],
        [
            ["特殊申请", "学生", "集中、分散、自主等实习方式和说明", "指导教师、管理员"],
            ["签到", "学生", "GPS 位置、时间、任务", "教师查看，异常按流程处理"],
            ["实习日志", "学生", "日期、标题、过程内容", "指导教师审核"],
            ["实习报告", "学生", "单位简介、目的、过程、总结和建议", "指导教师审核评分"],
            ["延期申请", "学生", "目标材料类型、申请期限和理由", "对应审核角色"],
            ["安全与保险", "学生或管理员", "承诺书、保单和有效期", "教师、管理员查看"],
        ],
        [2.5, 2.5, 7.5, 4.0],
    )
    add_note(document, "签到位置由 H5 调用设备定位获取，地图仅展示定位结果，不允许学生拖动修改。", "danger")

    add_heading(document, "5.8 审核、通过后修改与流程记录", 2)
    add_bullets(document, [
        "待审核记录显示通过和退回；已通过记录只显示通过后修改。",
        "退回或需修改后，学生完成修改并重新提交，状态恢复待审核。",
        "审核意见的最少字数、最多字数或不限制由模块配置控制，前后端同时校验。",
        "流程记录主线显示第几次提交、提交状态、内容和时间，支线显示审核角色、意见、评分和状态。",
        "通过后修改挂在当前审核支线下，学生下一次提交后形成新的提交主线。",
    ])

    add_heading(document, "5.9 成绩与归档", 2)
    add_bullets(document, [
        "教师按任务记录学生成绩，学生参与多个任务时各任务独立评分。",
        "课程成绩按计划设置的平均或累计等规则汇总。",
        "归档材料按计划、任务和学生查看，保险、安全承诺、日志、报告、大纲、指导书、实施表、教师报告和成绩分别展示。",
        "归档列表提供查看和导出，内容较长的材料必须进入详情查看。",
    ])

    add_heading(document, "6 实验实训管理", 1)
    add_heading(document, "6.1 一体化管理原则", 2)
    add_bullets(document, [
        "实验和实训共用一套业务框架，通过类别字段区分实验或实训。",
        "普通实验实训按年级管理，不使用毕业届次。",
        "教学计划、课表、项目、过程材料、成绩和归档均带类别，列表和权限按类别过滤。",
        "原实验和实训数据已迁移到统一结构，页面以“实验实训管理”作为主入口。",
    ])
    add_figure(document, "13-PC实验实训总览.png", "图 15  实验实训管理总览")

    add_heading(document, "6.2 教学计划与二维课表", 2)
    add_steps(document, [
        "新增教学计划，选择类别、年级、学院、专业、课程和负责教师。",
        "保存草稿或提交审核，审核通过后进入课表安排。",
        "课表按专业排课，横向为星期，纵向为第 1 至第 12 节。",
        "默认白天 8 节、晚上 4 节；每节开始和结束时间由课节配置维护。",
        "点击课表单元格安排课程、教师、场地和学生范围。",
        "发布前检查同一教师、场地和时间冲突。",
    ])
    add_note(document, "课节时间统一配置后影响后续排课显示，调整前应确认现有课表。", "warning")

    add_heading(document, "6.3 项目、过程与成绩", 2)
    add_bullets(document, [
        "课表安排确认后发布项目，并绑定参与班级和学生。",
        "学生通过 H5 完成签到、过程日志和总结报告。",
        "教师维护大纲、教案和成绩比例，审核过程材料并录入成绩。",
        "成绩比例用于计算考勤、项目实操、课程报告和总评成绩。",
        "反思报告提交审核后完成教学归档。",
        "场地管理维护校内场地或关联校外基地。",
    ])

    add_heading(document, "7 社会实践管理", 1)
    add_heading(document, "7.1 业务结构", 2)
    add_bullets(document, [
        "社会实践按年级、学院、专业和班级组织。",
        "管理范围包括总览、计划、集中实践、分散实践、指导教师、安全材料、签到与补签、成果材料、成绩、归档和统计。",
        "教务来源计划先同步到待确认区，确认后再生成社会实践计划。",
        "集中实践由学院统一建立项目、教师和学生范围；分散实践由学生申报，经审核后执行。",
    ])
    add_figure(document, "11-PC社会实践总览.png", "图 16  社会实践管理入口")

    add_heading(document, "7.2 计划与项目", 2)
    add_steps(document, [
        "管理员新增或同步社会实践计划，选择年级、学院、专业、实践方式和时间。",
        "保存草稿或提交审核；通过后建立集中或分散实践项目。",
        "集中项目绑定指导教师、班级和学生；分散项目审核学生申报。",
        "多人项目支持共同提交或分别提交，默认分别提交，具体由学院设置。",
    ])

    add_heading(document, "7.3 安全、过程、成果与成绩", 2)
    add_bullets(document, [
        "安全承诺书、保险材料和应急预案按项目或学生范围维护。",
        "家长知情书不是强制材料，由学院根据项目风险和年龄范围启用。",
        "签到与补签保留位置、时间、提交人和审核记录。",
        "成果材料包括实践报告、照片、证明等，按学院设置共同或分别提交。",
        "成绩规则区分集中和分散实践，可维护项目项、权重和满分。",
        "归档前检查计划、项目、指导教师、安全材料、过程、成果和成绩完整性。",
    ])

    add_heading(document, "8 消息中心与待办模板", 1)
    add_heading(document, "8.1 消息中心", 2)
    add_bullets(document, [
        "消息按日期和时间显示，分类为系统通知、待办提醒、处理结果和预警提醒。",
        "可按类型、已读状态和关键词查询，支持全部已读。",
        "带业务链接的消息可点击“打开关联页面”直接进入对应模块。",
        "有权限的管理员可向指定人员、角色或学校范围发送消息。",
    ])
    add_figure(document, "17-PC消息中心.png", "图 17  按时间分组的消息中心")

    add_heading(document, "8.2 流程待办和消息模板", 2)
    add_bullets(document, [
        "业务发送消息时只调用模板 code 并传递变量，不在业务代码中拼接正文。",
        "模板维护名称、code、消息类型、级别、标题、正文、跳转地址和启用状态。",
        "实习、实验实训、社会实践的提交待办、审核结果和通过后修改均预置模板。",
        "每周简报、保险到期、导出完成等定时或异步消息也使用模板。",
        "流程消息模板仅超级管理员可以维护。",
    ])
    add_figure(document, "18-PC消息模板管理.png", "图 18  流程待办和消息模板")

    add_heading(document, "9 文件、文档、模板、导出与收藏", 1)
    add_heading(document, "9.1 文件管理", 2)
    add_bullets(document, [
        "文件列表显示文件名、扩展名、MIME、分类、大小、上传人、上传时间、设备信息和状态。",
        "点击“打开”使用后端文件地址在新页面预览或下载。",
        "头像、壁纸、学校登录背景、材料模板和业务附件统一进入文件管理。",
        "临时文件由 is_temporary=1 和过期时间界定；到期后定时任务删除物理文件、软删除记录并更新引用计数。",
    ])
    add_figure(document, "19-PC文件管理.png", "图 19  文件管理列表")

    add_heading(document, "9.2 文档中心", 2)
    add_steps(document, [
        "维护文档分类，例如实践流程、学生帮助、教师帮助和管理员帮助。",
        "新增文档并填写标题、分类、版本、适用角色和发布状态。",
        "使用富文本编辑器维护正文，不直接填写 HTML 标签。",
        "发布后，授权角色可在 PC 或 H5 文档中心查看。",
        "修改内容时记录版本和变更说明。",
    ])
    add_figure(document, "20-PC文档中心.png", "图 20  文档中心列表和预览")
    add_figure(document, "21-PC富文本文档编辑.png", "图 21  富文本文档编辑")

    add_heading(document, "9.3 材料模板库", 2)
    add_bullets(document, [
        "模板库维护实习、实验实训、社会实践相关 Word、Excel、PDF 等材料。",
        "模板按分类、状态和关键词查询。",
        "学生和教师只查看、下载有权限且已发布的模板。",
        "计划导入等固定格式入口在操作区域提供对应模板下载。",
    ])

    add_heading(document, "9.4 导出任务", 2)
    add_bullets(document, [
        "列表导出创建异步任务，页面可继续操作。",
        "任务完成后在消息中心通知，并在导出任务中心提供下载。",
        "处理超过 30 分钟的任务由定时任务标记为超时，有权限者可重试。",
        "实施表、计划表、统计报表和归档材料按对应业务模板导出。",
    ])

    add_heading(document, "9.5 收藏夹", 2)
    add_bullets(document, [
        "新增收藏时填写标题、网址并上传图标。",
        "点击收藏项在新页面打开，避免覆盖当前业务窗口。",
        "收藏支持关键词搜索和加入桌面。",
    ])

    add_heading(document, "10 统计报表与日志审计", 1)
    add_heading(document, "10.1 统计报表", 2)
    add_bullets(document, [
        "实习报表包括总览、学院、专业、任务老师、学生过程和归档材料统计。",
        "实验实训提供成绩记载表，按教学计划汇总考勤、实操、报告和总分。",
        "筛选顺序为年级或毕业届次、学院、专业、班级、关键词。",
        "学院管理员默认本学院，专业管理员默认本专业，教师只统计本人任务学生。",
        "报表显示生成时间，列表支持分页和异步导出。",
    ])
    add_figure(document, "23-PC统计报表.png", "图 22  实习统计报表")

    add_heading(document, "10.2 日志审计", 2)
    add_bullets(document, [
        "所有 API 请求由全局中间件记录，不只记录登录。",
        "日志显示操作时间、用户、请求方法、接口、HTTP 状态、业务码、耗时、IP、响应消息和请求摘要。",
        "请求摘要包含接口用途，例如“查询角色权限”“获取实习筛选选项”，便于判断实际操作。",
        "日志按月分表、分页查询；用户管理中的日志仅查看单个用户记录。",
    ])
    add_figure(document, "22-PC日志审计.png", "图 23  接口用途和请求摘要日志")

    add_heading(document, "11 H5 移动端通用操作", 1)
    add_heading(document, "11.1 导航与页面交互", 2)
    add_bullets(document, [
        "底部导航根据当前角色和权限显示首页、实习和我的等入口。",
        "顶部消息按钮显示未读角标，点击进入消息中心。",
        "业务页使用顶部标签切换概况、任务、提交、审核和成绩。",
        "列表以卡片形式显示，筛选条件放入筛选面板，避免占用首屏。",
        "详情、审核和流程记录使用移动端弹层，底部按钮位于系统安全区和导航栏上方。",
        "页面转场使用轻量淡入和位移，遵循系统减少动态效果设置。",
    ])

    add_heading(document, "11.2 管理员移动端", 2)
    add_bullets(document, [
        "首页显示学校或组织范围内的任务、绑定、待审和异常概况。",
        "实习概况提供任务、待审和绑定统计；审核页处理待办，数据页查看授权范围列表。",
        "管理员 H5 与 PC 使用相同后端接口和数据权限。",
    ])
    add_mobile_pair(
        document,
        "25-H5管理员首页.png",
        "图 24  H5 管理员首页",
        "26-H5管理员实习概况.png",
        "图 25  H5 管理员实习概况",
    )

    add_heading(document, "11.3 教师移动端", 2)
    add_bullets(document, [
        "教师首页只显示本人任务的待评日志、待评报告和特殊申请。",
        "实习页按概况、审核和成绩切换，不显示学校级管理入口。",
        "审核列表按日志、报告、延期和特殊申请切换。",
        "待审核记录可通过或退回；已通过记录只显示通过后修改。",
    ])
    add_mobile_pair(
        document,
        "28-H5教师首页.png",
        "图 26  H5 教师首页",
        "29-H5教师审核列表.png",
        "图 27  H5 教师审核列表",
    )

    add_heading(document, "11.4 学生移动端", 2)
    add_bullets(document, [
        "学生首页显示本人任务、已评分任务、特殊申请和实习流程引导。",
        "任务页显示管理员分配的任务、课程、负责教师、开始结束时间和流程记录。",
        "提交页以卡片显示签到、日志、报告和延期申请，并展示阶段截止时间。",
        "日志或报告为需修改状态时显示“修改”按钮；修改后重新提交形成新的 recording。",
        "学生不需要选择指导教师，材料自动提交给任务绑定的负责教师。",
    ])
    add_mobile_pair(
        document,
        "31-H5学生首页.png",
        "图 28  H5 学生首页和流程引导",
        "32-H5学生任务与申请.png",
        "图 29  H5 学生任务和特殊申请",
    )
    add_mobile_pair(
        document,
        "33-H5学生提交入口.png",
        "图 30  H5 学生提交入口和截止时间",
        "34-H5学生日志与修改.png",
        "图 31  H5 学生日志和需修改操作",
    )

    add_heading(document, "11.5 查看流程记录", 2)
    add_steps(document, [
        "在列表卡片点击“记录”。",
        "查看第几次提交、提交状态、提交内容和时间。",
        "查看提交节点下的审核支线、审核意见、评分和状态。",
        "如存在通过后修改，继续查看学生下一次提交和后续审核。",
    ])
    add_mobile_pair(
        document,
        "30-H5流程记录.png",
        "图 32  H5 特殊申请流程记录",
        "35-H5学生日志流程记录.png",
        "图 33  H5 实习日志流程记录",
    )

    add_heading(document, "11.6 个人中心", 2)
    add_bullets(document, [
        "个人中心显示姓名、账号、当前角色、学校和数据范围。",
        "文档中心和模板库按当前角色展示可见内容。",
        "存在多个绑定账号时显示切换身份；只有一个账号时不显示。",
        "退出登录需要二次确认。",
    ])
    add_figure(document, "27-H5个人中心.png", "图 34  H5 个人中心", width=7.2)

    add_heading(document, "12 常见问题", 1)
    faq_rows = [
        ["登录后没有进入系统", "确认域名已绑定学校、接口地址使用当前域名、浏览器未拦截请求；再检查账号状态和角色权限。"],
        ["看不到某个模块", "检查菜单权限、模块显示状态和角色平台范围；教师和学生不会显示无权限的管理模块。"],
        ["选择学院后专业没有选项", "确认学院下已维护启用专业，并检查当前账号组织范围。"],
        ["筛选后仍显示全部数据", "确认当前年级或毕业届次、学院、专业默认值已加载，并检查账号组织范围配置。"],
        ["已通过记录不能再次审核", "属于正常规则。已通过记录只能按权限发起通过后修改。"],
        ["学生看不到修改按钮", "仅需修改状态显示修改按钮；待审核和已通过状态不允许学生修改。"],
        ["文件打开到系统页面", "应通过后端部署地址访问静态文件；开发前端端口不能直接代理后端文件路径。"],
        ["下拉框无法选择", "先确认上级年级或届次、学院、专业已选择，并刷新选项；仍异常时记录页面和账号提交管理员。"],
        ["导出没有立即下载", "导出为异步任务，请到消息中心或导出任务中心查看进度和下载结果。"],
        ["H5 定位失败", "检查浏览器定位权限、HTTPS、设备定位服务和网络；系统不允许手工拖动定位点。"],
        ["菜单改名后桌面未变化", "保存菜单后刷新权限或重新登录；桌面和启动台名称来自菜单配置。"],
        ["一键清理无法执行", "仅 APP_MODE=test 环境允许，正式环境固定禁止。"],
    ]
    add_table(document, ["问题", "处理方式"], faq_rows, [5.0, 11.5])

    add_heading(document, "13 日常维护顺序", 1)
    add_steps(document, [
        "维护学校基础配置、年级和毕业届次。",
        "维护学院、专业、班级和基地档案。",
        "新增或同步教师、学生、管理员账号，并配置角色。",
        "配置学院或专业管理员组织范围。",
        "检查菜单、按钮权限和桌面模块。",
        "维护课节时间、操作说明、文档和材料模板。",
        "维护消息模板和消息接收配置。",
        "创建或同步计划，完成任务拆分、审核和学生绑定。",
        "按日查看待办、异常签到、退回材料和消息发送情况。",
        "按周查看统计报表、导出任务、保险到期和系统日志。",
    ])

    add_heading(document, "14 附录", 1)
    add_heading(document, "14.1 关键权限边界", 2)
    add_table(
        document,
        ["操作", "超级/学校管理员", "学院管理员", "专业管理员", "教师", "学生"],
        [
            ["查看全校数据", "是", "否", "否", "否", "否"],
            ["维护用户与基础档案", "按权限", "授权范围", "授权范围", "否", "否"],
            ["维护实习计划和任务", "按权限", "授权学院", "授权专业", "执行本人任务", "否"],
            ["审核学生材料", "按节点", "按节点", "按节点", "本人任务", "否"],
            ["提交学生材料", "否", "否", "否", "否", "本人"],
            ["维护流程消息模板", "仅超级管理员", "否", "否", "否", "否"],
            ["一键登录", "按学校范围", "教师/学生", "教师/学生", "否", "否"],
        ],
        [3.2, 3.0, 2.8, 2.8, 2.2, 2.2],
    )

    add_heading(document, "14.2 安全要求", 2)
    add_bullets(document, [
        "正式环境必须使用 HTTPS，并正确配置学校域名。",
        "不要在登录页、公开文档或群聊中保存默认密码和一键登录链接。",
        "管理员操作应使用本人账号，禁止共享超级管理员账号。",
        "重要审核、任务变更、数据清理和同步生成均应执行二次确认。",
        "定期查看异常接口、登录失败、批量导出和权限变更日志。",
    ])

    add_note(document, "本手册中的菜单和按钮以当前账号权限为准。学校后续调整菜单名称、流程节点或操作说明后，应同步更新交付手册版本。", "success")

    document.core_properties.title = "高校实践管理系统操作手册"
    document.core_properties.subject = "成都锦城学院高校实践管理系统 PC/H5 完整交付操作手册"
    document.core_properties.author = "成都锦城学院"
    document.core_properties.keywords = "高校实践管理, 实习, 实验实训, 社会实践, 操作手册"
    document.core_properties.comments = "V2.0，基于 2026-08-11 当前本地交付版本"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_manual())
